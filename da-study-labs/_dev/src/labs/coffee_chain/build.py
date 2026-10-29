import random, sqlite3, json, datetime as dt, subprocess, os, shutil, sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "tools"))
from sitekit import recalc
os.makedirs(os.path.join(HERE, "out"), exist_ok=True)
os.chdir(os.path.join(HERE, "out"))   # all generated files go to out/


random.seed(2026)
stores = [(1,"Saigon Central","Ho Chi Minh City"),(2,"Hanoi Old Quarter","Hanoi"),(3,"Da Nang Beach","Da Nang")]
products = {"Latte":(55000,),"Americano":(45000,),"Cold Brew":(60000,),"Croissant":(40000,),"Matcha":(65000,)}
weights = {1:[5,3,4,2,2],2:[3,4,2,5,2],3:[2,2,5,3,4]}   # product popularity per store
cancel_p = {1:0.10,2:0.22,3:0.14}
rows=[]; oid=1001
days=[dt.date(2026,7,1)+dt.timedelta(d) for d in range(92)]
for i in range(90):
    d=random.choice(days)
    st=random.choice([1,1,2,2,3,3,1,2])
    pn=random.choices(list(products),weights[st])[0]
    qty=random.choice([1,1,2,2,3,4])
    status="cancelled" if random.random()<cancel_p[st] else "completed"
    rows.append([d,st,pn,qty,products[pn][0],status])
rows.sort(key=lambda r:(r[0],r[1]))
orders=[[1001+i]+r for i,r in enumerate(rows)]
cols=["order_id","order_date","store_id","product","qty","unit_price","status"]
df=pd.DataFrame(orders,columns=cols); df["order_date"]=pd.to_datetime(df["order_date"])
N=len(df)

# ---------- SQL (sqlite) ----------
con=sqlite3.connect(":memory:")
con.execute("CREATE TABLE stores(store_id INT, store_name TEXT, city TEXT)")
con.executemany("INSERT INTO stores VALUES(?,?,?)",stores)
con.execute("CREATE TABLE orders(order_id INT, order_date TEXT, store_id INT, product TEXT, qty INT, unit_price INT, status TEXT)")
con.executemany("INSERT INTO orders VALUES(?,?,?,?,?,?,?)",[(o[0],o[1].isoformat(),*o[2:]) for o in orders])

SQL={}
SQL[1]="""SELECT s.store_name,
       SUM(o.qty * o.unit_price) AS revenue
FROM orders o
JOIN stores s ON s.store_id = o.store_id
WHERE o.status = 'completed'
GROUP BY s.store_name
ORDER BY revenue DESC;"""
SQL[2]="""SELECT s.store_name,
       COUNT(*) AS total_orders,
       SUM(CASE WHEN o.status = 'cancelled' THEN 1 ELSE 0 END) AS cancelled,
       ROUND(100.0 * SUM(CASE WHEN o.status = 'cancelled' THEN 1 ELSE 0 END) / COUNT(*), 1) AS cancel_rate_pct
FROM orders o
JOIN stores s ON s.store_id = o.store_id
GROUP BY s.store_name
ORDER BY cancel_rate_pct DESC;"""
SQL[3]="""WITH monthly AS (
  SELECT strftime('%Y-%m', order_date) AS month,
         SUM(qty * unit_price) AS revenue
  FROM orders
  WHERE status = 'completed'
  GROUP BY 1
)
SELECT month,
       revenue,
       ROUND(100.0 * (revenue - LAG(revenue) OVER (ORDER BY month))
             / LAG(revenue) OVER (ORDER BY month), 1) AS mom_growth_pct
FROM monthly
ORDER BY month;"""
SQL[4]="""WITH by_product AS (
  SELECT store_id, product,
         SUM(qty * unit_price) AS revenue
  FROM orders
  WHERE status = 'completed'
  GROUP BY store_id, product
),
ranked AS (
  SELECT *,
         RANK() OVER (PARTITION BY store_id ORDER BY revenue DESC) AS rnk
  FROM by_product
)
SELECT s.store_name, r.product, r.revenue
FROM ranked r
JOIN stores s ON s.store_id = r.store_id
WHERE r.rnk = 1
ORDER BY s.store_name;"""
sql_res={k:con.execute(v).fetchall() for k,v in SQL.items()}

# ---------- pandas ----------
PY={}
PY[1]="""done = orders[orders["status"] == "completed"].copy()
done["revenue"] = done["qty"] * done["unit_price"]

result = (done.merge(stores, on="store_id")
              .groupby("store_name", as_index=False)["revenue"].sum()
              .sort_values("revenue", ascending=False))"""
PY[2]="""result = (orders.merge(stores, on="store_id")
               .assign(is_cancelled=lambda d: d["status"] == "cancelled")
               .groupby("store_name")
               .agg(total_orders=("order_id", "count"),
                    cancelled=("is_cancelled", "sum"))
               .reset_index())
result["cancel_rate_pct"] = (100 * result["cancelled"]
                             / result["total_orders"]).round(1)
result = result.sort_values("cancel_rate_pct", ascending=False)"""
PY[3]="""done = orders[orders["status"] == "completed"].copy()
done["revenue"] = done["qty"] * done["unit_price"]
done["month"] = done["order_date"].dt.strftime("%Y-%m")

result = done.groupby("month", as_index=False)["revenue"].sum()
result["mom_growth_pct"] = (result["revenue"].pct_change() * 100).round(1)"""
PY[4]="""done = orders[orders["status"] == "completed"].copy()
done["revenue"] = done["qty"] * done["unit_price"]

by_product = (done.groupby(["store_id", "product"], as_index=False)["revenue"].sum())
by_product["rnk"] = (by_product.groupby("store_id")["revenue"]
                               .rank(method="min", ascending=False))

result = (by_product[by_product["rnk"] == 1]
          .merge(stores, on="store_id")[["store_name", "product", "revenue"]]
          .sort_values("store_name"))"""
stores_df=pd.DataFrame(stores,columns=["store_id","store_name","city"])
py_res={}
for k,code in PY.items():
    env={"pd":pd,"orders":df.copy(),"stores":stores_df.copy()}
    exec(code,env)
    py_res[k]=env["result"]

def norm(rows): return [tuple(round(float(x),1) if isinstance(x,(int,float)) and not isinstance(x,bool) and x==x else x for x in r) for r in rows]
def nan2none(rows): return [tuple(None if (isinstance(x,float) and x!=x) else x for x in r) for r in rows]
for k in SQL:
    a=norm(nan2none(sql_res[k])); b=norm(nan2none(list(py_res[k].itertuples(index=False,name=None))))
    # sqlite window/first-row NULL vs pandas NaN handled by nan2none
    assert a==b,(k,a,b)
print("SQL == pandas for all 4 tasks")

# ---------- Excel ----------
wb=Workbook()
wo=wb.active; wo.title="orders"
hdr=cols+["line_total"]
wo.append(hdr)
for i,o in enumerate(orders, start=2):
    wo.append([o[0],o[1],o[2],o[3],o[4],o[5],o[6],f"=E{i}*F{i}"])
    wo.cell(i,2).number_format="yyyy-mm-dd"
last=N+1
ws=wb.create_sheet("stores"); ws.append(["store_id","store_name","city"])
for s in stores: ws.append(list(s))
O=lambda c: f"orders!${c}$2:${c}${last}"
XL={}
# Task 1
t1=wb.create_sheet("task1"); t1.append(["store_id","store_name","revenue"])
for i,s in enumerate(stores,start=2):
    t1.append([s[0],f"=INDEX(stores!$B$2:$B$4,MATCH(A{i},stores!$A$2:$A$4,0))",
               f'=SUMIFS({O("H")},{O("C")},A{i},{O("G")},"completed")'])
XL[1]=[("A2:A4","store_id typed or copied from the stores sheet",),
       ("B2","=INDEX(stores!$B$2:$B$4, MATCH(A2, stores!$A$2:$A$4, 0))"),
       ("C2",'=SUMIFS(orders!$H$2:$H$%d, orders!$C$2:$C$%d, A2, orders!$G$2:$G$%d, "completed")'%(last,last,last)),
       ("Sort","Data > Sort on column C, largest to smallest")]
# Task 2
t2=wb.create_sheet("task2"); t2.append(["store_id","store_name","total_orders","cancelled","cancel_rate_pct"])
for i,s in enumerate(stores,start=2):
    t2.append([s[0],f"=INDEX(stores!$B$2:$B$4,MATCH(A{i},stores!$A$2:$A$4,0))",
               f'=COUNTIFS({O("C")},A{i})',
               f'=COUNTIFS({O("C")},A{i},{O("G")},"cancelled")',
               f"=ROUND(100*D{i}/C{i},1)"])
XL[2]=[("C2","=COUNTIFS(orders!$C$2:$C$%d, A2)"%last),
       ("D2",'=COUNTIFS(orders!$C$2:$C$%d, A2, orders!$G$2:$G$%d, "cancelled")'%(last,last)),
       ("E2","=ROUND(100 * D2 / C2, 1)"),
       ("Sort","Data > Sort on column E, largest to smallest")]
# Task 3
t3=wb.create_sheet("task3"); t3.append(["month_start","month","revenue","mom_growth_pct"])
for i,m in enumerate([dt.date(2026,7,1),dt.date(2026,8,1),dt.date(2026,9,1)],start=2):
    t3.append([m,f'=TEXT(A{i},"yyyy-mm")',
               f'=SUMIFS({O("H")},{O("G")},"completed",{O("B")},">="&A{i},{O("B")},"<"&EDATE(A{i},1))',
               None if i==2 else f"=ROUND(100*(C{i}/C{i-1}-1),1)"])
    t3.cell(i,1).number_format="yyyy-mm-dd"
XL[3]=[("A2:A4","1 Jul 2026, 1 Aug 2026, 1 Sep 2026 (first day of each month)"),
       ("B2",'=TEXT(A2, "yyyy-mm")'),
       ("C2",'=SUMIFS(orders!$H$2:$H$%d, orders!$G$2:$G$%d, "completed", orders!$B$2:$B$%d, ">="&A2, orders!$B$2:$B$%d, "<"&EDATE(A2, 1))'%(last,last,last,last)),
       ("D3","=ROUND(100 * (C3 / C2 - 1), 1)")]
# Task 4: matrix store x product + pick max
plist=list(products)
t4=wb.create_sheet("task4")
t4.append(["store_id","store_name"]+plist+["top_product","top_revenue"])
for i,s in enumerate(stores,start=2):
    row=[s[0],f"=INDEX(stores!$B$2:$B$4,MATCH(A{i},stores!$A$2:$A$4,0))"]
    for j,p in enumerate(plist):
        col=get_column_letter(3+j)
        row.append(f'=SUMIFS({O("H")},{O("C")},$A{i},{O("D")},{col}$1,{O("G")},"completed")')
    row.append(f"=INDEX($C$1:$G$1,MATCH(MAX(C{i}:G{i}),C{i}:G{i},0))")
    row.append(f"=MAX(C{i}:G{i})")
    t4.append(row)
XL[4]=[("C1:G1","Product names typed as column headers"),
       ("C2",'=SUMIFS(orders!$H$2:$H$%d, orders!$C$2:$C$%d, $A2, orders!$D$2:$D$%d, C$1, orders!$G$2:$G$%d, "completed")'%(last,last,last,last)),
       ("H2","=INDEX($C$1:$G$1, MATCH(MAX(C2:G2), C2:G2, 0))"),
       ("I2","=MAX(C2:G2)")]
for sh in wb.worksheets:
    for c in sh[1]: c.font=Font(bold=True); c.fill=PatternFill("solid",fgColor="DDE6EF")
    for i in range(1,sh.max_column+1): sh.column_dimensions[get_column_letter(i)].width=16
wb.save("lab.xlsx")
RECALC = recalc('lab.xlsx', "recalc")
if RECALC:
    from openpyxl import load_workbook
    rv=load_workbook("recalc/lab.xlsx",data_only=True)
    def vals(name,r0,r1,c0,c1):
        sh=rv[name]; return [[sh.cell(r,c).value for c in range(c0,c1+1)] for r in range(r0,r1+1)]
    x1={r[1]:r[2] for r in vals("task1",2,4,1,3)}
    assert x1=={r[0]:r[1] for r in sql_res[1]},(x1,sql_res[1])
    x2={r[1]:(r[2],r[3],r[4]) for r in vals("task2",2,4,1,5)}
    assert x2=={r[0]:(r[1],r[2],r[3]) for r in sql_res[2]},(x2,sql_res[2])
    x3=vals("task3",2,4,2,4)
    for a,b in zip(x3,sql_res[3]):
        assert a[0]==b[0] and a[1]==b[1] and (a[2] is None and b[2] is None or abs(a[2]-b[2])<0.05),(a,b)
    x4={r[1]:(r[7],r[8]) for r in vals("task4",2,4,1,9)}
    assert x4=={r[0]:(r[1],r[2]) for r in sql_res[4]},(x4,sql_res[4])
    print("Excel (LibreOffice recalculated) == SQL for all 4 tasks")

# tie check for task 4 (RANK would keep ties, INDEX/MATCH picks first)
for sid in (1,2,3):
    d=df[(df.store_id==sid)&(df.status=="completed")].assign(r=lambda d:d.qty*d.unit_price).groupby("product").r.sum().sort_values(ascending=False)
    assert d.iloc[0]!=d.iloc[1],("tie",sid)

# ---------- export ----------
def res_table(k):
    heads={1:["store_name","revenue"],2:["store_name","total_orders","cancelled","cancel_rate_pct"],3:["month","revenue","mom_growth_pct"],4:["store_name","product","revenue"]}[k]
    return {"cols":heads,"rows":[list(r) for r in sql_res[k]]}
data={
 "orders":[[o[0],o[1].isoformat(),o[2],o[3],o[4],o[5],o[6]] for o in orders],
 "stores":[list(s) for s in stores],
 "sql":SQL,"py":PY,
 "xl":{k:[list(x) for x in v] for k,v in XL.items()},
 "res":{k:res_table(k) for k in SQL},
}
json.dump(data,open("data.json","w"),ensure_ascii=False)
shutil.copy(RECALC or "lab.xlsx","coffee-chain-lab.xlsx")
for k in SQL: print(k,res_table(k))
print(N,"orders; completed",int((df.status=="completed").sum()))
